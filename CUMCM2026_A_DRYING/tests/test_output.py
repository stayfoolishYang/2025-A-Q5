"""TEST_ONLY reconstruction, event and ephemeral workbook checks.

Synthetic trajectories in this module are never installed in results/runs or
reported as real drying calculations. They stress the public module contracts.
"""
from dataclasses import replace
from decimal import Decimal
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

from drying.inputs import Inputs
from drying.spatial import Grid,DryingSystem
from drying.manufactured import ManufacturedCase
from drying.physics import T0,C0,ELL,surface
from drying.reconstruction import reconstruct,base_nodes,ReconstructionError,OutsideDomainError
from drying.events import affine_strict_interval,upward_report_time,scan_events,verify_report
from drying.export import regular_times,output_row_count,export_workbooks,verify_workbook
from drying.trajectory import fingerprint

RAW=Path(__file__).resolve().parents[1]/"data"/"raw"


@pytest.fixture(scope="module")
def inputs():return Inputs(RAW)


class SyntheticTrajectory:
    def __init__(self,times,states):
        self.times=np.asarray(times,float);self.states=np.asarray(states,float)
        self.start_time=float(times[0]);self.end_time=float(times[-1])
        self.index={"identity":{"run_id":"TEST_ONLY_OUTPUT_FIXTURE","config":"TEST_ONLY"},"times":list(map(float,times))}
    def iter_segments(self):
        for i in range(len(self.times)-1):
            yield self.times[i],self.states[i],self.times[i+1],self.states[i+1]
    def at(self,t):
        if t<self.start_time or t>self.end_time:raise ValueError("OUT_OF_COVERAGE")
        i=min(max(int(np.searchsorted(self.times,t,side="right")-1),0),len(self.times)-2)
        w=(t-self.times[i])/(self.times[i+1]-self.times[i])
        return (1-w)*self.states[i]+w*self.states[i+1]


def uniform(system,c=1.,t=310.):
    return np.tile([t,c],system.grid.ncell)


def test_o1_axis_from_exact_quadratic_means(inputs):
    g=Grid(7);s=DryingSystem(g,inputs,question=23,h=0,hm=0)
    m=(g.xi_faces[:-1]**2+g.xi_faces[1:]**2)/2
    y=np.stack([310-2*m,2-.4*m],axis=-1).reshape(-1)
    r=reconstruct(s,100,y)
    assert r.query(0)==pytest.approx((310,2),abs=2e-13)
    assert r.max_C==pytest.approx(2,abs=5e-16)
    assert r.max_C>y.reshape(-1,2)[:,1].max()
    assert r.query(.02)==r.surface()
    assert r.query(.02)[1]==y[-1]
    with pytest.raises(OutsideDomainError):r.query(.02000000001)


def test_initial_state_distinct_from_right_limit_flux(inputs):
    s=DryingSystem(Grid(5),inputs,question=1)
    a=reconstruct(s,0,s.initial())
    b=reconstruct(s,0,s.initial(),side="right")
    assert a.surface()==(T0,C0)
    assert b.surface()[1]<C0
    assert a.initial_semantics=="INITIAL_STATE"


def test_full_c_nodes_tensor_convexity_and_surface_shared(inputs):
    g=Grid(5,4,"C");s=DryingSystem(g,inputs,question=23,end_condition="C1")
    mr=(g.xi_faces[:-1]**2+g.xi_faces[1:]**2)/2
    mz=(g.z_faces[1:]**3-g.z_faces[:-1]**3)/(3*g.dz*ELL**2)
    p=(1-mr[:,None])*(1-mz[None,:])
    y=np.stack((310+2*p,1+.2*p),axis=-1).reshape(-1)
    r=reconstruct(s,100,y)
    assert r.query(0,0)==pytest.approx((312,1.2),abs=2e-13)
    assert r.nodes.shape==(7,6,2)
    ev=s.evaluate(100,y,jacobian=False)
    for j,z in enumerate(g.z):
        assert r.surface(z)==pytest.approx(ev.surface_values[j],abs=2e-13)
    for i,x in enumerate(g.xi):
        assert r.query(.02*x,ELL)==pytest.approx(ev.end_values[i],abs=2e-13)
    rng=np.random.default_rng(11)
    for x,z in rng.random((100,2)):
        v=r.query(.02*x,ELL*z)[1]
        assert r.nodes[...,1].min()-1e-15<=v<=r.max_C+1e-15


def test_mms_uses_constant_material_and_nonuniform_environment(inputs):
    g=Grid(6,4,"C");case=ManufacturedCase(route="C",moving=True)
    s=DryingSystem(g,inputs,question=4,geometry="moving",end_condition="C1",test_case=case)
    r=reconstruct(s,200,case.averages(g,200))
    assert r.query(0,0)==pytest.approx(case.exact(0,0,200),abs=2e-13)
    assert np.ptp(r.nodes[-1,:,1])>0


def test_reconstruction_invalid_axes_are_not_clipped(inputs):
    s=DryingSystem(Grid(3),inputs)
    y=uniform(s,2.55);y[1]=2.55;y[3]=1.0
    with pytest.raises(ReconstructionError,match="envelope"):
        reconstruct(s,1,y)
    y=uniform(s,1.);y[1]=.01;y[3]=2.0
    with pytest.raises(ReconstructionError,match="nonpositive"):
        reconstruct(s,1,y)


def test_exact_domain_at_measured_radius_and_crossing(inputs):
    s=DryingSystem(Grid(4),inputs,question=4,geometry="moving",h=0,hm=0)
    r=reconstruct(s,259200,uniform(s))
    assert r.domain_relation(.01198)==0
    assert r.query(.01198)==r.surface()
    assert r.domain_relation(float(Decimal(".01198000000000001")))==1
    assert r.domain_relation(float(Decimal(".01197999999999999")))==-1
    before=reconstruct(s,0,uniform(s));after=reconstruct(s,259200,uniform(s))
    assert before.domain_relation(.015)<0 and after.domain_relation(.015)>0


def test_affine_strict_interval_interior_gap_and_strict_equalities():
    x=affine_strict_interval([.16,.12,.10],[.12,.16,.10])
    assert x.lower==pytest.approx(.25) and x.upper==pytest.approx(.75)
    assert x.lower_open and x.upper_open
    assert affine_strict_interval([.15],[.15]) is None
    assert affine_strict_interval([.149],[.149]).lower==0
    assert affine_strict_interval([.16,.14],[.14,.16]) is None
    assert affine_strict_interval([.15],[.14]).lower_open
    assert affine_strict_interval([.14],[.15]).upper_open
    with pytest.raises(ValueError):affine_strict_interval([np.nan],[.1])


def test_quadratic_true_curve_counterexample_not_claimed_affine():
    true=lambda x:(x-.3)*(x-.4)
    assert all(true(t)>0 for t in [0,.5,1]) and true(.35)<0
    assert affine_strict_interval([true(0)],[true(1)],threshold=0) is None


def test_scanner_candidate_rebound_and_earlier_gap(inputs):
    s=DryingSystem(Grid(3),inputs,h=0,hm=0)
    tr=SyntheticTrajectory([0,10,20],[uniform(s,.16),uniform(s,.14),uniform(s,.13)])
    a=scan_events(s,tr)
    assert a["status"]=="PROVISIONAL_EVENT" and a["t_hat"]==pytest.approx(5)
    assert a["earliest_verified"] and a["retention_verified"]
    tr2=SyntheticTrajectory([0,10,20],[uniform(s,.16),uniform(s,.14),uniform(s,.16)])
    b=scan_events(s,tr2)
    assert b["status"]=="EVENT_UNRESOLVED" and b["rebound"]
    tr3=SyntheticTrajectory([5,10,20],[uniform(s,.16),uniform(s,.14),uniform(s,.13)])
    c=scan_events(s,tr3)
    assert not c["earliest_verified"]


def test_report_grid_and_output_unique_final_times():
    assert upward_report_time(1.)==1.08
    assert upward_report_time(1.08)==1.08
    assert list(regular_times(120,60))==[60,120]
    assert list(regular_times(120.24,60))==[60,120,120.24]
    assert list(regular_times(.36,60))==[.36]
    assert output_row_count(1800,1)==1800


def valid_record(tr,t=120.24):
    nominal=119.99
    plus=119.984
    minus=119.996
    return {"status":"DRYING_COMPLETE","earliest_verified":True,"retention_verified":True,
            "t_report":t,"eta_strict":.0002,"t_hat":nominal,"tL":119.986,"tR":119.994,
            "t_plus":plus,"t_minus":minus,"t_eligible":120.,"et_ref":.01,
            "level_time_uncertainty":max(nominal-plus,minus-nominal),"report_delay":t-nominal,
            "error_evidence":{"eC":.0001,"eG":.0001,"eT":.01,"et":.05,"coverage_end":tr.end_time,
                              "checks_passed":True,"refinement_verified":True,"run_ids":["TEST_ONLY_FIXTURE"],
                              "config_fingerprint":"TEST_ONLY","trajectory_fingerprint":fingerprint(tr.index)}}


def test_report_requires_actual_margin_and_covering_identity(inputs):
    s=DryingSystem(Grid(3),inputs,h=0,hm=0)
    tr=SyntheticTrajectory([0,120.24],[uniform(s,.16),uniform(s,.14)])
    rec=valid_record(tr)
    assert verify_report(s,tr,rec)["strict_sum"]==pytest.approx(.1402)
    for change in [{"eta_strict":1e-12},{"t_report":120.25},{"status":"PROVISIONAL_EVENT"},
                   {"earliest_verified":False},{"t_report":120.60,"report_delay":20.6}]:
        with pytest.raises(ValueError):verify_report(s,tr,{**rec,**change})
    bad={**rec,"error_evidence":{**rec["error_evidence"],"trajectory_fingerprint":"WRONG"}}
    with pytest.raises(ValueError):verify_report(s,tr,bad)
    wrong_config={**rec,"error_evidence":{**rec["error_evidence"],"config_fingerprint":"OTHER_CONFIGURATION"}}
    with pytest.raises(ValueError,match="identity"):
        verify_report(s,tr,wrong_config)


def test_q4_export_mask_last_row_and_candidate_readback(inputs,tmp_path):
    s=DryingSystem(Grid(3),inputs,question=4,geometry="moving",h=0,hm=0)
    # Synthetic low-field final values test serialization only in pytest tmp.
    tr=SyntheticTrajectory([0,259200],[uniform(s,.14),uniform(s,.14)])
    rec=valid_record(tr,259200.)
    nominal,plus,minus=259199.984,259199.974,259199.994
    rec.update(t_hat=nominal,tL=259199.980,tR=259199.988,t_plus=plus,t_minus=minus,
               t_eligible=259199.999,report_delay=259200.-nominal,
               level_time_uncertainty=max(nominal-plus,minus-nominal))
    out=export_workbooks(s,tr,tmp_path/"test_only_export",4,rec)
    assert len(out["files"])==2
    assert out["reports"]["q4_domain_mask"]["all_rows_checked"]==4320
    w=load_workbook(Path(out["files"][0]),read_only=True,data_only=True)
    try:
        last=None
        for row in w["Sheet1"].iter_rows(values_only=True):last=row
        assert last[0]==259200 and len(last)==23
        assert last[13] is None and last[-1]==pytest.approx(.14)
    finally:w.close()
    with pytest.raises(FileExistsError):export_workbooks(s,tr,tmp_path/"test_only_export",4,rec)


def test_q23_export_same_source_and_no_nominal_export(inputs,tmp_path):
    s=DryingSystem(Grid(3),inputs,question=23,h=0,hm=0)
    tr=SyntheticTrajectory([0,120.24],[uniform(s,.16),uniform(s,.14)])
    with pytest.raises(ValueError):export_workbooks(s,tr,tmp_path/"bad",23,{"status":"PROVISIONAL_EVENT"})
    out=export_workbooks(s,tr,tmp_path/"test_only_export",23,valid_record(tr))
    assert out["reports"]["q23_same_source"]["common_times_checked"]==3
    assert out["reports"]["result2.xlsx"]["rows_per_sheet"]==121


def test_q1_full_window_guard_and_roundtrip(inputs,tmp_path):
    s=DryingSystem(Grid(3),inputs,question=1,h=0,hm=0)
    tr=SyntheticTrajectory([0,1800],[s.initial(),s.initial()])
    with pytest.raises(ValueError):export_workbooks(s,tr,tmp_path/"bad",1)
    out=export_workbooks(s,tr,tmp_path/"test_only_export",1,developer_approved=True)
    assert out["reports"]["result1.xlsx"]["rows_per_sheet"]==1800
    assert out["reports"]["result1.xlsx"]["max_sample_error"]<1e-13


def test_display_01500_is_not_threshold_equality(inputs):
    s=DryingSystem(Grid(3),inputs,h=0,hm=0)
    tr=SyntheticTrajectory([0,120.24],[uniform(s,.16),uniform(s,.1499999)])
    rec=valid_record(tr)
    rec["eta_strict"]=2e-8
    rec["error_evidence"].update(eC=1e-8,eG=1e-8)
    check=verify_report(s,tr,rec)
    assert format(check["max_C"],".4f")=="0.1500"
    assert check["strict_sum"]<.15
