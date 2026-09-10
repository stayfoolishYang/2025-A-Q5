"""TEST_ONLY: isolated symbolic/algebraic Stage05 checks; no PDE/ODE march.

Reads original inputs and installed dependency metadata. Appends actual results
only to stage05_presolve_checks.json. Never creates solver/results workbooks.
"""
from pathlib import Path
from fractions import Fraction as Q
from decimal import Decimal
from datetime import datetime, timezone
import hashlib
import importlib.metadata as metadata
import json
import sys

import sympy as sy
from openpyxl import load_workbook
from scipy.integrate import BDF
from scipy.sparse.linalg import splu


HERE = Path(__file__).resolve().parent
EVIDENCE = HERE / "stage05_presolve_checks.json"
record = json.loads(EVIDENCE.read_text(encoding="utf-8"))
checks = []


def check(name, condition, observed, scope="TEST_ONLY_EXECUTED"):
    assert bool(condition), (name, observed)
    checks.append({"id": name, "status": "PASS", "observed": observed,
                   "evidence_class": scope})


# General step-ratio identities at t_new=0, with exact symbolic histories.
h, ratio = sy.symbols("h ratio", positive=True)
times = [sy.Integer(0), -h, -h-h/ratio]
alpha = [(1+2*ratio)/(1+ratio), -(1+ratio), ratio**2/(1+ratio)]
check("BDF_EQUAL_STEP", [sy.simplify(x.subs(ratio, 1)) for x in alpha]
      == [sy.Rational(3, 2), -2, sy.Rational(1, 2)], "[3/2,-2,1/2]")
for degree in range(3):
    actual = sy.simplify(sum(a*t**degree for a, t in zip(alpha, times))/h)
    expected = 1 if degree == 1 else 0
    check(f"BDF_POLYNOMIAL_{degree}", actual == expected, str(actual))
parasitic = ratio**2/(1+2*ratio)
check("ZERO_EQUATION_INCREMENT_RATIO", sy.simplify(parasitic.subs(ratio, sy.Rational(3, 2)))
      == sy.Rational(9, 16), "q<=1.5 gives increment multiplier <=9/16; not nonlinear stability proof")

# Isolated analytic scalar y'=lambda*y expressions, not a time integrator.
lam = sy.symbols("lambda", real=True)
y_bdf = (2-sy.exp(-lam*h)/2)/(sy.Rational(3, 2)-lam*h)
y_be = 1/(1-lam*h)
y_half_half = 1/(1-lam*h/2)**2
lead = sy.simplify(sy.limit((y_bdf-y_be)/h**2, h, 0))
be_est = sy.simplify(sy.limit(2*(y_be-y_half_half)/h**2, h, 0))
check("BDF_BE_MONITOR_LEADING_ORDER", lead == -lam**2/2, str(lead)+" times h^2")
check("BE_STEP_DOUBLING_MONITOR", be_est == lam**2/2, str(be_est)+" times h^2")

# Current-state mass matrix contribution; normalized and raw Newton forms.
c, temp, hist = sy.symbols("C T history", real=True)
ac = 2+c
forcing = temp*c+3
S = sy.Rational(3, 2)*temp+hist
normalized = S-h*forcing/ac
raw = ac*S-h*forcing
identity = sy.simplify(sy.diff(raw,c)/ac-sy.diff(normalized,c)-sy.diff(ac,c)*normalized/ac)
check("MASS_JACOBIAN_EQUIVALENCE", identity == 0, str(identity))
omission = sy.simplify(sy.diff(normalized,c)-(-h*sy.diff(forcing,c)/ac))
check("MASS_JACOBIAN_OMISSION_DETECTABLE", omission != 0, str(omission))

# Series-resistance identities and limiting cases.
A,b,delta,un,ue = sy.symbols("A b delta UN Ue", positive=True)
us=ue+A/(A+b*delta)*(un-ue)
ps=b*(us-ue)
check("SURFACE_FLUX_IDENTITY", sy.simplify(A*(un-us)/delta-ps)==0, "0")
check("SURFACE_B_ZERO", sy.simplify(us.subs(b,0)-un)==0 and ps.subs(b,0)==0, "Us=UN; p=0")
check("SURFACE_EQUILIBRIUM", sy.simplify(ps.subs(un,ue))==0, "p=0")
check("SURFACE_ZERO_DISTANCE_LIMIT", sy.limit(us,delta,0)==un, "Us -> UN")
for coeff,exchange,distance in [(Q(3,7),Q(1,5),Q(2,9)),(Q(1,10**8),Q(8,10**7),Q(1,1000))]:
    beta=coeff/(coeff+exchange*distance)
    check("SURFACE_CONVEX_COEFFICIENT",0<=beta<=1,str(beta))

# Telescoping with unrelated exact face values. No state or field evolution.
edges=[Q(0),Q(1,3),Q(2,3),Q(1)]
omega=[(edges[i+1]**2-edges[i]**2)/2 for i in range(3)]
flux=[Q(0),Q(2,11),Q(-3,13),Q(5,17)]
R=Q(1,50)
water_sum=sum(-2*(edges[i+1]*flux[i+1]-edges[i]*flux[i])/R for i in range(3))
check("RADIAL_WATER_TELESCOPE",water_sum==-2*flux[-1]/R,str(water_sum+2*flux[-1]/R))
ell=Q(1,8); dz=ell/2
radial=[[Q(0),Q(1,7),Q(-2,9),Q(4,11)], [Q(0),Q(-1,13),Q(2,15),Q(3,17)]]
axial=[[Q(0),Q(i+1,19),Q(i+2,23)] for i in range(3)]
total=sum((Q(2)/ell)*(-dz/R*(edges[i+1]*radial[j][i+1]-edges[i]*radial[j][i])-omega[i]*(axial[i][j+1]-axial[i][j])) for i in range(3) for j in range(2))
boundary= -Q(2)/(ell*R)*sum(dz*row[-1] for row in radial)-Q(2)/ell*sum(omega[i]*axial[i][-1] for i in range(3))
check("HALF_DOMAIN_WATER_TELESCOPE",total==boundary,str(total-boundary))

# Newly repaired MMS interface: independent source-weighted telescoping.
sources_b=[Q(2,7),Q(-3,11),Q(5,13)]
rhs_b=[-(edges[i+1]*flux[i+1]-edges[i]*flux[i])/R+omega[i]*sources_b[i] for i in range(3)]
water_with_source=sum(2*omega[i]*(rhs_b[i]/omega[i]) for i in range(3))
source_rate_b=sum(2*omega[i]*sources_b[i] for i in range(3))
check("MMS_B_DISCRETE_SOURCE_BALANCE",water_with_source+2*flux[-1]/R-source_rate_b==0,
      "0; weights 2*omega; static algebra only")
sources_c=[[Q(i+j+1,29) for j in range(2)] for i in range(3)]
source_rate_c=sum(Q(2)*omega[i]*dz/ell*sources_c[i][j] for i in range(3) for j in range(2))
rhs_c=[[-dz/R*(edges[i+1]*radial[j][i+1]-edges[i]*radial[j][i])-omega[i]*(axial[i][j+1]-axial[i][j])+omega[i]*dz*sources_c[i][j] for j in range(2)] for i in range(3)]
water_with_source_c=sum(Q(2)*omega[i]*dz/ell*(rhs_c[i][j]/(omega[i]*dz)) for i in range(3) for j in range(2))
check("MMS_C_DISCRETE_SOURCE_BALANCE",water_with_source_c-boundary-source_rate_c==0,
      "0; weights 2*omega*dz/ell; static algebra only")

# Symbolically verify MMS, including explicitly averaged 2D initial field.
x,z,t,tau,rr,length = sy.symbols("xi z t tau R ell", positive=True)
amp,base,diffus,capacity = sy.symbols("amp base diffus capacity", positive=True)
E=sy.exp(-t/tau); p=1-x*x; q=1-(z/length)**2
for dimension in (1,2):
    axial_shape=1 if dimension==1 else q
    u=base+amp*E*p*axial_shape
    lap=sy.diff(x*sy.diff(u,x),x)/(rr**2*x)
    if dimension==2: lap+=sy.diff(u,z,2)
    source=amp*E*(-capacity*p*axial_shape/tau+4*diffus*axial_shape/rr**2+(0 if dimension==1 else 2*diffus*p/length**2))
    check(f"MMS_{dimension}D_SOURCE",sy.simplify(capacity*sy.diff(u,t)-diffus*lap-source)==0,"0; symbolic, not solved")
    side=sy.simplify(-diffus*sy.diff(u,x).subs(x,1)/rr)
    check(f"MMS_{dimension}D_SIDE",sy.simplify(side-2*diffus*amp*E*axial_shape/rr)==0,"0")
    def normalized_average(expr):
        radial_average=sy.integrate(2*x*expr,(x,0,1))
        return radial_average if dimension==1 else sy.integrate(radial_average,(z,0,length))/length
    boundary_rate=2*side/rr
    if dimension==2:
        boundary_rate=sy.integrate(boundary_rate,(z,0,length))/length
        end=-diffus*sy.diff(u,z).subs(z,length)
        boundary_rate+=sy.integrate(2*x*end,(x,0,1))/length
    source_balance=sy.simplify(normalized_average(capacity*sy.diff(u,t))+boundary_rate-normalized_average(source))
    check(f"MMS_{dimension}D_GLOBAL_SOURCE_BALANCE",source_balance==0,
          "0; exact weighted integral of capacity*u_t + outward flux - volume source")
    radius_test=rr*(1-t/(10*tau))
    volume_test=2*sy.pi*radius_test**2*length
    storage_test=volume_test*capacity*normalized_average(u)
    reynolds_term=sy.diff(volume_test,t)*capacity*normalized_average(u)
    moving_heat_balance=sy.simplify(sy.diff(storage_test,t)-reynolds_term+volume_test*(boundary_rate-normalized_average(source)).subs(rr,radius_test))
    check(f"MMS_{dimension}D_MOVING_HEAT_SOURCE_BALANCE",moving_heat_balance==0,
          "0; derivative of volume*capacity*u minus explicit Reynolds volume term + outward flux - source")
avg=sy.integrate(sy.integrate(x*(base+amp*p*q),(z,0,length/2)),(x,0,sy.Rational(1,2)))/(sy.Rational(1,8)*length/2)
expected=base+amp*(1-sy.Rational(1,8))*(1-sy.Rational(1,12))
check("MMS_2D_INITIAL_AVERAGE",sy.simplify(avg-expected)==0,str(sy.simplify(avg)))
wrong=(base+amp*(1-sy.Rational(1,8)))*(1-sy.Rational(1,12))
check("MMS_BASE_MULTIPLICATION_COUNTEREXAMPLE",sy.simplify(wrong-avg)==-base/12,"wrong whole-expression multiplication shifts base by -base/12")

# Independent heat diagnostic must detect a deliberately inconsistent path.
path_integral=sy.integrate((1-t/10)**2,(t,0,1))
check("HEAT_PATH_DIAGNOSTIC_NOT_IDENTICALLY_ZERO",path_integral!=0,str(path_integral)+" for T'=1, V=(1-t/10)^2, a=1, boundary flux=0")

# Isolated event functions; no material simulation or drying-time result.
theta=sy.symbols("theta",real=True)
g=(theta-sy.Rational(3,10))*(theta-sy.Rational(4,10))
sampled=[g.subs(theta,v) for v in (0,sy.Rational(1,2),1)]
check("QUADRATIC_EVENT_FINITE_SAMPLES_MISS",all(v>0 for v in sampled) and g.subs(theta,sy.Rational(7,20))<0,
      {"samples_0_half_1":list(map(str,sampled)),"roots":["3/10","2/5"],"negative_interval":"(0.30,0.40)","scope":"not represented by endpoint-linear interpolation"})
lower=Q(0); upper=Q(1)
for value,slope in [(Q(16,100),Q(-4,100)),(Q(12,100),Q(4,100))]:
    crossing=(Q(15,100)-value)/slope
    if slope<0: lower=max(lower,crossing)
    else: upper=min(upper,crossing)
check("LINEAR_BASE_NODE_INTERVAL_INTERSECTION",(lower,upper)==(Q(1,4),Q(3,4)),"(1/4,3/4), open")
check("SLOW_CROSSING_ERROR_SCALE",Q(1,100000)/Q(1,10**8)==1000,"TEST_ONLY: 1e-5 kg/kg / 1e-8 kg/kg/s = 1000 s, unrelated to 0.01 s root resolution")

# Pure timestamp sets: fractional final row, no duplicate, no extrapolation.
for report in [Q(9009,25),Q(360)]:
    for step in (1,60):
        rows=[Q(i*step) for i in range(1,int(report//step)+1)]
        if not rows or rows[-1]!=report: rows.append(report)
        check(f"FINAL_ROW_{report}_{step}",rows[-1]==report and len(rows)==len(set(rows)) and all(v<=report for v in rows),
              {"toy_report_s":str(report),"step_s":step,"rows_with_header":len(rows)+1})
eligible=Q(259200)+Q(1,100); quantum=Q(9,25)
rounded=quantum*(-(-eligible//quantum))
check("REPORT_GRID_OUTSIDE_HORIZON_REJECT",rounded>Q(259200),{"toy_candidate":str(eligible),"report_grid":str(rounded),"required_status":"OUT_OF_COVERAGE_NOT_COMPLETE"})

# Sources remain read-only. Recheck true schemas relevant to final-row contract.
source_results=[]
for source in record['source_hashes']:
    source_path=Path(source['path']); digest=hashlib.sha256(source_path.read_bytes()).hexdigest()
    check("SOURCE_HASH_"+source_path.name,digest==source['sha256'],digest,"SPEC_RECHECKED")
    if source_path.suffix.lower()=='.xlsx':
        book=load_workbook(source_path,read_only=True,data_only=True)
        source_results.append({'file':source_path.name,'sheets':[{'name':ws.title,'rows':ws.max_row,'cols':ws.max_column,'header':[cell for cell in next(ws.iter_rows(min_row=1,max_row=1,values_only=True))]} for ws in book.worksheets]})
        if source_path.name=='附件1.xlsx':
            values=list(book.active.iter_rows(min_row=2,values_only=True))
            ymin=min(Decimal(str(row[2])) for row in values); ymax=max(Decimal(str(row[2])) for row in values)
            check('ENVIRONMENT_BELOW_THRESHOLD',ymax<Decimal('0.15'),{'min':str(ymin),'max':str(ymax),'count':len(values)},'SPEC_RECHECKED')
        book.close()

doc=BDF.__doc__
check('INSTALLED_BDF_IDENTITY',all(x in doc for x in ['1 to 5','quasi-constant','NDF']),
      {'scipy':metadata.version('scipy'),'claims':['variable order 1 to 5','quasi-constant steps','NDF enhancement'],'custom_BDF2_equivalent':False},'SPEC_RECHECKED')
check('SPARSE_LU_CAPABILITY',callable(splu),'imported scipy.sparse.linalg.splu; no linear solve or PDE run','SPEC_RECHECKED')
record['test_only_execution']={'timestamp':datetime.now(timezone.utc).isoformat(),'runtime':sys.executable,'python':sys.version,'packages':{name:metadata.version(name) for name in ['numpy','scipy','sympy','openpyxl']},'checks':checks,'assertions_passed':len(checks),'script':str(Path(__file__).resolve()),'script_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),'source_schema_recheck':source_results,'formal_PDE':'NOT_RUN','formal_convergence':'NOT_RUN','formal_workbooks':'NOT_RUN'}
temp_path=EVIDENCE.with_suffix('.json.tmp'); temp_path.write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n',encoding='utf-8'); json.loads(temp_path.read_text(encoding='utf-8')); temp_path.replace(EVIDENCE)
print(json.dumps({'assertions_passed':len(checks),'evidence':str(EVIDENCE),'scope':'TEST_ONLY_EXECUTED; FORMAL_NUMERICAL_VALIDATION_NOT_RUN'},ensure_ascii=False))
