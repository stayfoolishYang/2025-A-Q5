"""Verified attachment reconstruction; A26-04-v2 sections 5.2--5.4.

No extrapolation is delegated to library defaults. Environment observations
are Celsius, radius observations are centimetres; public output is SI.
"""
from dataclasses import dataclass
from hashlib import sha256
import json
import math
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.interpolate import PchipInterpolator


SOURCE_HASHES = {
    "environment.xlsx": "7ef32870abeef420b89560b2530ff60dfe4255917805151d89988d0311af9dd7",
    "radius.xlsx": "5563acbfa4b4afb10cc6c03e2207e5369bf39da27576672aff14cc5c32e704af",
    "templates/result1.xlsx": "23b261b295c1b787d000eebbca6521c37075107b6fcf78724f8d395ce1798ff4",
    "templates/result2.xlsx": "23b261b295c1b787d000eebbca6521c37075107b6fcf78724f8d395ce1798ff4",
    "templates/result3.xlsx": "07e4793d620a7f899804c0298d49a16a197960440fd47f8bb780c57ec27e2859",
    "templates/result4.xlsx": "86e9300ffa3d30c43de895ea6723da943e85b8740b137bcae5af7107f076eeac",
}


class InputError(ValueError):
    pass


class InputHorizonError(InputError):
    pass


@dataclass(frozen=True)
class Boundary:
    T_inf: float
    C_eq: float
    R: float
    Rdot: float


def _table(path, headers, rows):
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        if wb.sheetnames != ["Sheet1"]:
            raise InputError(f"INPUT_INVALID: {path.name}: expected Sheet1")
        ws = wb["Sheet1"]
        values = list(ws.iter_rows(values_only=True))
        if len(values) != rows + 1 or len(values[0]) != len(headers) or tuple(values[0]) != headers:
            raise InputError(f"INPUT_INVALID: {path.name}: schema/header mismatch")
        if any(len(r) != len(headers) or any(isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) for v in r) for r in values[1:]):
            raise InputError(f"INPUT_INVALID: {path.name}: missing/nonfinite/formula/non-numeric value")
        return np.asarray(values[1:], dtype=np.float64)
    finally:
        wb.close()


class Inputs:
    def __init__(self, data_dir, scenario="S0", radius_method="linear", radius_tail="NONE", window_start_h=3.0):
        self.data_dir = Path(data_dir)
        if scenario not in {"S0", "S1"} or radius_method not in {"linear", "pchip"} or radius_tail not in {"NONE", "HOLD"}:
            raise InputError("INPUT_INVALID: scenario/radius_method/radius_tail")
        if window_start_h not in {2.5, 3.0, 3.5} or (scenario == "S1" and window_start_h != 3.0):
            raise InputError("INPUT_INVALID: unapproved environment window")
        self.scenario, self.radius_method, self.radius_tail = scenario, radius_method, radius_tail
        self.window_start_h = float(window_start_h)
        try:
            self.manifest = json.loads((self.data_dir / "source_manifest.json").read_text(encoding="utf-8-sig"))
        except (OSError, ValueError) as exc:
            raise InputError("INPUT_INVALID: missing or malformed source_manifest.json") from exc
        listed = {v["relative_path"]: v["sha256"] for v in self.manifest.get("files", [])}
        self.hashes = {}
        for relative, expected in SOURCE_HASHES.items():
            p = self.data_dir / relative
            try:
                actual = sha256(p.read_bytes()).hexdigest()
            except OSError as exc:
                raise InputError(f"INPUT_INVALID: missing source {relative}") from exc
            if actual != expected or listed.get(relative) != expected:
                raise InputError(f"INPUT_INVALID: source hash mismatch {relative}")
            self.hashes[relative] = actual
        e = _table(self.data_dir / "environment.xlsx", ("时间", "温度", "水分浓度"), 241)
        r = _table(self.data_dir / "radius.xlsx", ("时间", "半径"), 145)
        if not np.array_equal(e[:, 0], np.arange(241) * 60.0):
            raise InputError("INPUT_INVALID: environment time index")
        if not np.array_equal(r[:, 0], np.arange(145) * 1800.0):
            raise InputError("INPUT_INVALID: radius time index")
        if np.any(e[:, 1] + 273.15 <= 0) or np.any(e[:, 2] < 0):
            raise InputError("INPUT_INVALID: environment units/domain")
        if np.any(r[:, 1] <= 0) or np.any(np.diff(r[:, 1]) > 0):
            raise InputError("INPUT_INVALID: radius must be positive/nonincreasing")
        self.env_times, self.temp_c, self.water = e.T.copy()
        self.radius_times = r[:, 0].copy()
        self.radius_cm = r[:, 1].copy()
        self.radius_m = r[:, 1].copy() * 0.01
        mask = self.env_times >= self.window_start_h * 3600
        self.window_count = int(mask.sum())
        self.tail_temp_c = float(self.temp_c[mask].mean()) if scenario == "S0" else float(self.temp_c[-1])
        self.tail_water = float(self.water[mask].mean()) if scenario == "S0" else float(self.water[-1])
        self._pchip = PchipInterpolator(self.radius_times, self.radius_m, extrapolate=False) if radius_method == "pchip" else None
        if self._pchip is not None:
            self._validate_pchip()
        for a in (self.env_times, self.temp_c, self.water, self.radius_times, self.radius_cm, self.radius_m):
            a.setflags(write=False)
        payload = {"hashes": self.hashes, "scenario": scenario, "radius_method": radius_method,
                   "radius_tail": radius_tail, "window_start_h": window_start_h, "input_version": "A26-INPUT-v1"}
        self.fingerprint = sha256(json.dumps(payload, sort_keys=True, separators=(",", ":")).encode()).hexdigest()

    def _validate_pchip(self):
        # Inspect each polynomial at endpoints, stationary values and the
        # derivative's vertex. Thus positivity/monotonicity are not sample-only.
        tol = 64 * np.finfo(float).eps
        for i, h in enumerate(np.diff(self.radius_times)):
            a, b, c, d = self._pchip.c[:, i]
            candidates = [0., h]
            roots = np.roots([3*a, 2*b, c]) if a != 0 else ([-c/(2*b)] if b != 0 else [])
            candidates.extend(float(x.real) for x in roots if abs(x.imag if np.iscomplexobj(x) else 0) < tol and 0 < x.real < h)
            vals = [((a*x+b)*x+c)*x+d for x in candidates]
            dcandidates = [0., h]
            if a != 0 and 0 < -b/(3*a) < h:
                dcandidates.append(-b/(3*a))
            derivatives = [(3*a*x+2*b)*x+c for x in dcandidates]
            if min(vals) <= 0 or max(derivatives) > tol:
                raise InputError(f"INPUT_INVALID: pchip segment {i} violates positive nonincreasing radius")

    @staticmethod
    def _check(t, question, geometry, side):
        if not np.isscalar(t) or not math.isfinite(float(t)) or t < 0:
            raise InputError("INPUT_INVALID: time must be finite and nonnegative")
        if question not in {1, 23, 4} or geometry not in {"fixed", "moving"} or side not in {"point", "left", "right"}:
            raise InputError("INPUT_INVALID: question/geometry/side")

    def at(self, t, question=23, geometry="fixed", side="point"):
        self._check(t, question, geometry, side)
        t = float(t)
        tail = t > 14400 or (t == 14400 and side == "right")
        tc = self.tail_temp_c if tail else float(np.interp(t, self.env_times, self.temp_c))
        ce = self.tail_water if tail else float(np.interp(t, self.env_times, self.water))
        radius, rate = 0.02, 0.0
        if geometry == "moving":
            if t > 259200:
                if self.radius_tail != "HOLD":
                    raise InputHorizonError("INPUT_HORIZON_REACHED: radius coverage ends at 259200 s")
                radius = float(self.radius_m[-1])
            elif self.radius_tail == "HOLD" and t == 259200 and side == "right":
                radius = float(self.radius_m[-1])
            elif self._pchip is not None:
                radius, rate = float(self._pchip(t)), float(self._pchip.derivative()(t))
            else:
                # Point/left uses the preceding segment at a knot; right uses
                # the succeeding segment, with first/last endpoints protected.
                ix = int(np.searchsorted(self.radius_times, t, side="right" if side == "right" else "left") - 1)
                ix = min(max(ix, 0), len(self.radius_times)-2)
                h = self.radius_times[ix+1] - self.radius_times[ix]
                rate = float((self.radius_m[ix+1]-self.radius_m[ix])/h)
                radius = float(self.radius_m[ix] + (t-self.radius_times[ix])*rate)
        return Boundary(tc + 273.15, ce, radius, rate)

    def nodes(self, question=23, geometry="fixed", tmax=259200):
        self._check(tmax, question, geometry, "point")
        n = self.env_times[1:]
        # PCHIP is C1, but retained observation nodes are deliberately aligned
        # under the authority's conservative all-input-node restart policy.
        if geometry == "moving":
            n = np.union1d(n, self.radius_times[1:])
        return n[n <= tmax].copy()

    def node_kind(self, t, question=23, geometry="fixed"):
        self._check(t, question, geometry, "point")
        if t == 14400 and self.scenario == "S0":
            return "JUMP"
        if t > 0 and (np.any(self.env_times == t) or (geometry == "moving" and np.any(self.radius_times == t))):
            return "KINK"
        return "NONE"
