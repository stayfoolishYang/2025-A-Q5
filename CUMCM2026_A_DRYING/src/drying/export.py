"""Streaming template-compatible candidate workbooks with readback first.

Openpyxl is the existing verified server dependency. No workbook formulas or
simulated values are synthesized by this module: all rows query one trajectory.
"""
from copy import copy
from decimal import Decimal
import csv
import hashlib
import json
import math
import os
from pathlib import Path

import numpy as np
from openpyxl import Workbook,load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter

from .events import verify_report
from .reconstruction import reconstruct,OutsideDomainError
from .trajectory import fingerprint

RADII_CM = tuple(float(Decimal(i)/10) for i in range(21))
RADII_M = tuple(float(Decimal(i)/1000) for i in range(21))


def regular_times(end,step):
    """Exact decimal/rational set union of regular times and unique last row."""
    e=Decimal(str(end)); d=Decimal(str(step))
    if not e.is_finite() or e<=0 or not d.is_finite() or d<=0:
        raise ValueError("invalid output time range")
    n=int(e//d)
    for k in range(1,n+1):
        yield float(k*d)
    if e!=n*d:
        yield float(e)


def output_row_count(end,step):
    e,d=Decimal(str(end)),Decimal(str(step))
    if not e.is_finite() or e<=0 or not d.is_finite() or d<=0:
        raise ValueError("invalid output time range")
    n=int(e//d)
    return n+int(e!=n*d)


def _cell(ws,value,template_cell,header=False):
    c=WriteOnlyCell(ws,value=value)
    cache=getattr(ws,"_drying_style_cache",None)
    if cache is None:
        cache={};ws._drying_style_cache=cache
    key=(template_cell.coordinate,header)
    if key not in cache:
        c.font=copy(template_cell.font); c.fill=copy(template_cell.fill)
        c.border=copy(template_cell.border); c.alignment=copy(template_cell.alignment)
        c.protection=copy(template_cell.protection)
        c.number_format=template_cell.number_format if header else "0.0000"
        cache[key]=c._style
    else:
        c._style=cache[key]
    return c


def _create_sheet(wb,template,columns):
    ws=wb.create_sheet(template.title)
    ws.sheet_format=copy(template.sheet_format)
    ws.sheet_properties=copy(template.sheet_properties)
    ws.page_margins=copy(template.page_margins)
    ws.page_setup=copy(template.page_setup)
    ws.print_options=copy(template.print_options)
    ws.freeze_panes=template.freeze_panes
    for col in range(1,columns+1):
        source=min(col,template.max_column)
        dim=template.column_dimensions.get(get_column_letter(source))
        if dim is not None:
            ws.column_dimensions[get_column_letter(col)]=copy(dim)
            ws.column_dimensions[get_column_letter(col)].index=get_column_letter(col)
    if template.row_dimensions[1].height:
        ws.row_dimensions[1].height=template.row_dimensions[1].height
    return ws


def _row_values(rec,question):
    values=[]; mask=[]
    for radius in RADII_M:
        try:
            values.append(rec.query(radius));mask.append(True)
        except OutsideDomainError:
            if question!=4:
                raise
            values.append((None,None));mask.append(False)
    if question==4:
        values.append(rec.surface())
    return values,mask


def verify_workbook(path,system,trajectory,question,end):
    """Full structure/value-domain check plus independent trajectory samples."""
    step=1 if question in (1,2) else 60
    expected_names=["温度","水分浓度"] if question in (1,2) else ["Sheet1"]
    ncols=23 if question==4 else 22
    wb=load_workbook(path,read_only=True,data_only=False)
    rows_expected=output_row_count(end,step)
    # Query first/last/middle and all paper-prescribed times within coverage.
    prescribed=([100,300,600,900,1200,1500,1800] if question==1 else
                [1800,3600,5400,7200,9000,10800] if question==2 else
                list(range(21600,int(end)+1,21600)))
    samples={1,max(1,rows_expected//2),rows_expected}
    sample_times=set(float(t) for t in prescribed if t<=end)
    report={"path":str(path),"status":"PASS","structure":"ALL_ROWS",
            "numerical_verification":"TRAJECTORY_SAMPLES_PLUS_FULL_Q4_DOMAIN_MASK",
            "rows_per_sheet":rows_expected,"columns":ncols,"sheets":{},"max_sample_error":0.}
    try:
        if wb.sheetnames!=expected_names:
            raise ValueError("OUTPUT_INVALID: sheet names")
        for name in expected_names:
            ws=wb[name]; rows=ws.iter_rows(values_only=True)
            header=next(rows,None)
            expected_header=["时间\\到药材中心的距离",*RADII_CM]
            if question==4: expected_header.append("药材表面")
            if tuple(header or ())!=tuple(expected_header):
                raise ValueError("OUTPUT_INVALID: expanded header")
            if ws.cell(2,2).number_format!="0.0000":
                raise ValueError("OUTPUT_INVALID: four-decimal display format")
            expected=iter(regular_times(end,step));count=0;sample_count=0
            for row in rows:
                count+=1
                t=next(expected,None)
                if t is None or len(row)!=ncols or row[0]!=t:
                    raise ValueError(f"OUTPUT_INVALID: row {count+1} time/column count")
                if any(v is not None and (isinstance(v,bool) or not isinstance(v,(float,int)) or not math.isfinite(v)) for v in row):
                    raise ValueError("OUTPUT_INVALID: text, nonfinite or Excel error value")
                sample=count in samples or t in sample_times
                rec=reconstruct(system,t,trajectory.at(t)) if question==4 or sample else None
                eq=0 if name=="温度" else 1
                for j,v in enumerate(row[1:]):
                    outside=question==4 and j<21 and rec.domain_relation(RADII_M[j])>0
                    if outside != (v is None):
                        raise ValueError(f"OUTPUT_INVALID: row {count+1}, column {j+2} missingness")
                    if v is not None and eq==1 and v<=0:
                        raise ValueError("OUTPUT_INVALID: nonpositive moisture")
                    if sample and not outside:
                        ref=(rec.surface() if question==4 and j==21 else rec.query(RADII_M[j]))[eq]
                        if eq==0: ref-=273.15
                        err=abs(v-ref);report["max_sample_error"]=max(report["max_sample_error"],err)
                        if err>2e-12*max(1.,abs(ref)):
                            raise ValueError("OUTPUT_INVALID: numerical sample mismatch")
                if question==4:
                    if abs(row[-1]-rec.surface()[1])>2e-12:
                        raise ValueError("OUTPUT_INVALID: surface column")
                if sample:sample_count+=1
            if count!=rows_expected or next(expected,None) is not None:
                raise ValueError("OUTPUT_INVALID: truncated rows")
            report["sheets"][name]={"checked_rows":count,"numerical_sample_rows":sample_count}
    finally:
        wb.close()
    report["sha256"]=hashlib.sha256(Path(path).read_bytes()).hexdigest()
    return report


def _verify_q23_same_source(path2,path3):
    a=load_workbook(path2,read_only=True,data_only=True)
    b=load_workbook(path3,read_only=True,data_only=True)
    checked=0
    try:
        aa=a["水分浓度"].iter_rows(min_row=2,values_only=True)
        current=next(aa,None)
        for row in b["Sheet1"].iter_rows(min_row=2,values_only=True):
            while current is not None and current[0]<row[0]:current=next(aa,None)
            if current is None or current[0]!=row[0] or any(abs(x-y)>2e-12 for x,y in zip(current[1:],row[1:])):
                raise ValueError("OUTPUT_INVALID: Q2/Q3 common-time mismatch")
            checked+=1
    finally:
        a.close();b.close()
    return {"status":"PASS","common_times_checked":checked}


def export_workbooks(system,trajectory,outdir,question,event_record=None,developer_approved=False):
    if question not in (1,23,4) or question!=system.question:
        raise ValueError("OUTPUT_INVALID: question/system mismatch")
    if system.test_case is not None:
        raise ValueError("OUTPUT_INVALID: TEST_CASE cannot produce formal result workbooks")
    if trajectory.start_time!=0:
        raise ValueError("OUTPUT_INVALID: complete source must begin at initial state")
    if question==1:
        if not developer_approved or trajectory.end_time<1800:
            raise ValueError("OUTPUT_UNAPPROVED: Q1 requires developer acceptance and 1800 s coverage")
        end=1800.;strict=None
    else:
        strict=verify_report(system,trajectory,event_record or {})
        end=strict["t_report"]
    root=Path(outdir).resolve();raw=system.inputs.data_dir.resolve()
    if root==raw or raw in root.parents:
        raise ValueError("OUTPUT_INVALID: cannot overwrite source area")
    root.mkdir(parents=True,exist_ok=True)
    questions=[2,3] if question==23 else [question]
    finals=[root/f"result{q}.xlsx" for q in questions]
    mask_final=root/"q4_domain_mask.csv" if question==4 else None
    manifest_final=root/"candidate_manifest.json"
    for p in [*finals,manifest_final,*([mask_final] if mask_final else [])]:
        if p.exists():raise FileExistsError(p)
    for q in questions:
        rows=output_row_count(end,1 if q in (1,2) else 60)
        if rows>min(1048575,getattr(system,"max_output_rows",1048575)):
            raise ValueError("OUTPUT_CAPACITY_EXCEEDED")
    staged=[];reports={};mask_rows=0
    for q,final in zip(questions,finals):
        temporary=final.with_name(final.stem+".tmp.xlsx")
        if temporary.exists():raise FileExistsError(temporary)
        template=load_workbook(raw/"templates"/f"result{q}.xlsx",data_only=False)
        wb=Workbook(write_only=True)
        sheets=[]
        ncols=23 if q==4 else 22
        try:
            for template_sheet in template:
                ws=_create_sheet(wb,template_sheet,ncols)
                header=[template_sheet.cell(1,1).value,*RADII_CM]
                if q==4:header.append("药材表面")
                ws.append([_cell(ws,value,template_sheet.cell(1,min(i+1,6)),True) for i,value in enumerate(header)])
                sheets.append((ws,template_sheet))
            mask_tmp=mask_final.with_suffix(".tmp.csv") if mask_final else None
            if mask_tmp is not None and mask_tmp.exists():raise FileExistsError(mask_tmp)
            mask_file=mask_tmp.open("w",encoding="utf-8",newline="") if mask_tmp else None
            try:
                mask_writer=csv.writer(mask_file) if mask_file else None
                if mask_writer:
                    mask_writer.writerow(["t_s","R_m",*[f"r_{v:g}_cm_inside" for v in RADII_CM],"max_C","max_xi","max_z_m"])
                for t in regular_times(end,1 if q in (1,2) else 60):
                    rec=reconstruct(system,t,trajectory.at(t))
                    values,mask=_row_values(rec,q)
                    for ws,ts in sheets:
                        eq=0 if ws.title=="温度" else 1
                        numbers=[t]+[v[eq]-(273.15 if eq==0 else 0.) if v[eq] is not None else None for v in values]
                        ws.append([_cell(ws,value,ts.cell(2,min(i+1,6))) for i,value in enumerate(numbers)])
                    if mask_writer:
                        mask_writer.writerow([t,rec.R,*[int(v) for v in mask],rec.max_C,*rec.max_position]);mask_rows+=1
                wb.save(temporary)
            finally:
                if mask_file:mask_file.close()
            reports[f"result{q}.xlsx"]=verify_workbook(temporary,system,trajectory,q,end)
            reports[f"result{q}.xlsx"]["path"]=str(final)
            staged.append((temporary,final))
        finally:
            template.close()
    if question==23:
        reports["q23_same_source"]=_verify_q23_same_source(staged[0][0],staged[1][0])
    if question==4:
        with mask_tmp.open(encoding="utf-8",newline="") as f:
            reader=csv.reader(f);next(reader);rows=0
            expected=iter(regular_times(end,60))
            for row in reader:
                rows+=1;t=next(expected,None)
                if t is None or float(row[0])!=t or len(row)!=26:
                    raise ValueError("OUTPUT_INVALID: Q4 mask structure")
                rec=reconstruct(system,t,trajectory.at(t))
                if abs(float(row[1])-rec.R)>1e-15 or [int(x) for x in row[2:23]]!=[int(rec.domain_relation(r)<=0) for r in RADII_M]:
                    raise ValueError("OUTPUT_INVALID: Q4 mask/source mismatch")
            if rows!=mask_rows or next(expected,None) is not None:raise ValueError("OUTPUT_INVALID: Q4 mask rows")
        staged.append((mask_tmp,mask_final))
        reports["q4_domain_mask"]={"status":"PASS","all_rows_checked":mask_rows}
    manifest={"status":"STAGE06_CANDIDATE","stage07":"NOT_RUN","model_version":"A26-04-v2",
              "input_fingerprint":system.inputs.fingerprint,"trajectory_identity":trajectory.index["identity"],
              "trajectory_fingerprint":fingerprint(trajectory.index),"question":question,"end_time_s":end,
              "route":system.grid.route,"slice_z_m":0,"strict_report":strict,"event_record":event_record,
              "developer_approved":bool(developer_approved),"reports":reports,"files":[p.name for _,p in staged],
              "verification_scope":"Full structure and Q4 masks/surface; sampled numerical cells. Not independent Stage07 validation."}
    temp_manifest=manifest_final.with_suffix(".tmp.json")
    if temp_manifest.exists():raise FileExistsError(temp_manifest)
    temp_manifest.write_text(json.dumps(manifest,ensure_ascii=False,indent=2,allow_nan=False)+"\n",encoding="utf-8")
    json.loads(temp_manifest.read_text(encoding="utf-8"))
    # All files have passed before any candidate filename is published. The
    # manifest is the final group commit marker; interruptions before it are
    # distinguishable from a complete candidate bundle.
    for temp,final in staged:
        os.replace(temp,final)
    os.replace(temp_manifest,manifest_final)
    return {"files":[str(p) for _,p in staged],"manifest":str(manifest_final),"reports":reports}
