"""Independent reopen check of every delivered numeric cell, time, and Q4 mask."""
from pathlib import Path
import json
import openpyxl
import numpy as np

ROOT=Path(__file__).resolve().parents[1]

def main():
    result=[]
    payload=json.loads((ROOT/'results'/'workbooks.json').read_text(encoding='utf-8'))
    for book in payload:
        q=book['q'];file=ROOT/'results'/f'result{q}.xlsx'
        wb=openpyxl.load_workbook(file,data_only=True,read_only=True)
        assert wb.sheetnames==[s['name'] for s in book['sheets']]
        checked=0
        for spec in book['sheets']:
            ws=wb[spec['name']];actual=list(ws.values);expected=spec['rows']
            assert len(actual)==len(expected) and max(map(len,actual))==len(expected[0])
            for a,b in zip(actual,expected):
                for x,y in zip(a,b):
                    if isinstance(y,(float,int)):
                        assert isinstance(x,(float,int)) and abs(x-y)<1e-10,(q,x,y)
                        checked+=1
                    else:assert x==y,(q,x,y)
            times=np.array([r[0] for r in actual[1:]])
            assert np.all(np.diff(times)>0)
            if q<=2:
                assert np.array_equal(times,np.arange(1,1801 if q==1 else 10801))
            else:
                assert np.array_equal(times[:-1],np.arange(60,times[-1],60))
            assert all(v is None or v>=0 for row in actual[1:] for v in row[1:])
        meta=json.loads((ROOT/'results'/f'q{q}.json').read_text())
        source=np.load(ROOT/'results'/f'q{q}.npz')['data'];n=meta['n']
        assert np.max(source[:,2+n:]-source[:,2+n,None])<1e-9
        if q>=3:
            assert np.max(source[-1,2+n:])<=.15
            assert meta['event_bracket_s'][1]-meta['event_bracket_s'][0]<=.1
        if q==4:
            expected_r=np.loadtxt(ROOT/'results'/'q4_radius_output.csv',delimiter=',',skiprows=1)
            for row,(_,r) in zip(book['sheets'][0]['rows'][1:],expected_r):
                for pos,v in zip(np.arange(21)/10,row[1:-1]):
                    assert (v is None)==(pos>r+1e-10)
        result.append(dict(file=file.name,sheets=wb.sheetnames,rows=len(source)-1,numeric_cells_checked=checked,
                           maximum_final_C=meta['final_max_C'],status='passed'))
        wb.close()
    (ROOT/'results'/'export_validation.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps(result,indent=2))

if __name__=='__main__':main()
