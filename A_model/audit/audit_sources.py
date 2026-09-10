"""Read original XLSX/PDF and compare official archive hashes; no input edits."""
from pathlib import Path
import hashlib
import json
import os
import zipfile
import openpyxl
from pypdf import PdfReader

dest=Path(__file__).resolve().parent
manifest=json.loads((dest.parent/'data/manifest.json').read_text(encoding='utf-8'))
findings={}
for name, source in manifest.items():
    p=Path(source['source'])
    record={'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
    if p.suffix=='.xlsx':
        wb=openpyxl.load_workbook(p,data_only=False)
        record['defined_names']=list(wb.defined_names)
        record['external_links']=len(wb._external_links)
        record['sheets']=[]
        for ws in wb:
            record['sheets'].append(dict(title=ws.title,state=ws.sheet_state,rows=ws.max_row,columns=ws.max_column,
                hidden_rows=[k for k,v in ws.row_dimensions.items() if v.hidden],
                hidden_columns=[k for k,v in ws.column_dimensions.items() if v.hidden],
                formulas=[(c.coordinate,c.value) for row in ws for c in row if c.data_type=='f'],
                comments=[(c.coordinate,c.comment.text) for row in ws for c in row if c.comment],
                first_rows=list(ws.values)[:4],last_rows=list(ws.values)[-5:]))
        with zipfile.ZipFile(p) as z:
            record['archive_members']=z.namelist()
            record['metadata_xml']={n:z.read(n).decode('utf-8') for n in z.namelist()
                 if n=='docProps/custom.xml' or n.endswith('.rels')}
    else:
        pdf=PdfReader(p)
        record['page_count']=len(pdf.pages)
        record['boundary_relevant_excerpts']=['预热平衡和恒温干燥两个阶段',
            '附件1给出了烘干初期各时间点（单位：s）烘房的温度（单位：°C）和水分浓度（单位：kg/kg）']
        record['embedded_attachment_names']=list(pdf.attachments)
    findings[name]=record
archive=Path(os.environ['TEMP'])/'CUMCM2026Problems-audit.zip'
with zipfile.ZipFile(archive) as z:
    matched={}
    for info in z.infolist():
        if info.is_dir(): continue
        digest=hashlib.sha256(z.read(info)).hexdigest()
        for name,source in manifest.items():
            if digest==source['sha256']: matched[name]={'archive_member':info.filename,'sha256':digest}
    findings['official_archive']={'url':'https://www.mcm.edu.cn/upload_cn/CUMCM2026Problems.zip',
        'download_date':'2026-09-10','sha256':hashlib.sha256(archive.read_bytes()).hexdigest(),
        'matching_local_sources':matched,'members':z.namelist()}
(dest/'source_evidence.json').write_text(json.dumps(findings,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
print(json.dumps(findings['official_archive'],ensure_ascii=False,indent=2))
